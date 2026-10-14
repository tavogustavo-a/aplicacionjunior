"""Exporta listado Trailero (tanqueos) a TXT, Excel y PDF."""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


from trailero_pdf import format_fecha_slash


def _pdf_cell(text: str, style: ParagraphStyle) -> Paragraph:
    safe = (
        str(text or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )
    return Paragraph(safe, style)


def _money(n: float | int | None) -> str:
    try:
        return f"${float(n or 0):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except (TypeError, ValueError):
        return "$0,00"


def _num(n: float | int | None, decimals: int = 3) -> str:
    try:
        if n is None:
            return ""
        fmt = f"{{:.{decimals}f}}"
        return fmt.format(float(n)).replace(".", ",")
    except (TypeError, ValueError):
        return ""


def _meta(payload: dict) -> str:
    hora = payload.get("hora_desde") or "00:00"
    return (
        f"{payload.get('imap_usuario', '')} · "
        f"{payload.get('carpeta', '')} · "
        f"desde {hora} · "
        f"{format_fecha_slash(payload.get('desde', ''))} → "
        f"{format_fecha_slash(payload.get('hasta', ''))}"
    )


def _rows_ok(payload: dict) -> list[dict]:
    return [d for d in (payload.get("detalle") or []) if d.get("ok")]


def _rows_bad(payload: dict) -> list[dict]:
    return [d for d in (payload.get("detalle") or []) if not d.get("ok")]


def export_txt(payload: dict) -> bytes:
    lines = [
        "APLICACIONJUNIOR · SERVICIO TRAILERO · TANQUEOS",
        _meta(payload),
        f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        "",
        f"Tanqueos OK: {payload.get('tanqueos_total', 0)}",
        f"No leídos / error: {payload.get('pdfs_error', 0)}",
        f"TOTAL: {_money(payload.get('total_general'))}",
        "",
        "=== LISTADO (por fecha) ===",
        f"{'#':<5} {'Fecha':<12} {'Placa':<10} {'Cantidad':>10} "
        f"{'P.Galon':>12} {'Dcto':>12} {'Valor Total':>14}",
        "-" * 80,
    ]
    for d in _rows_ok(payload):
        lines.append(
            f"{str(d.get('tanqueo') or ''):<5} "
            f"{format_fecha_slash(d.get('fecha') or '')[:12]:<12} "
            f"{(d.get('placa') or '')[:10]:<10} "
            f"{_num(d.get('cantidad')):>10} "
            f"{_money(d.get('precio_galon')):>12} "
            f"{_money(d.get('descuento')):>12} "
            f"{_money(d.get('valor_total')):>14}"
        )
    bad = _rows_bad(payload)
    if bad:
        lines.append("")
        lines.append("=== NO LEÍDOS / ERROR ===")
        for d in bad:
            lines.append(
                f"- {(d.get('archivo') or d.get('asunto') or 'correo')}: "
                f"{d.get('error') or 'Error'}"
            )
    lines.append("")
    return ("\n".join(lines) + "\n").encode("utf-8")


def export_xlsx(payload: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Acpm"

    headers = [
        "Tanqueo #",
        "Fecha",
        "Placa",
        "CANTIDAD",
        "Precio x Galon",
        "Descuento",
        "Valor Total",
    ]
    header_fill = PatternFill("solid", fgColor="4F46E5")
    header_font = Font(color="FFFFFF", bold=True)
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_i, d in enumerate(_rows_ok(payload), start=2):
        fecha = format_fecha_slash(d.get("fecha") or "")
        ws.cell(row_i, 1, d.get("tanqueo"))
        ws.cell(row_i, 2, fecha)
        ws.cell(row_i, 3, d.get("placa") or "")
        ws.cell(row_i, 4, float(d.get("cantidad") or 0))
        ws.cell(row_i, 5, float(d.get("precio_galon") or 0))
        ws.cell(row_i, 6, float(d.get("descuento") or 0))
        ws.cell(row_i, 7, float(d.get("valor_total") or 0))
        for col in (5, 6, 7):
            ws.cell(row_i, col).number_format = '#,##0.00'

    # Hoja de errores si hay
    bad = _rows_bad(payload)
    if bad:
        ws_e = wb.create_sheet("No leidos")
        ws_e.append(["Archivo / Asunto", "Error", "Factura", "Fecha", "Placa"])
        for d in bad:
            ws_e.append(
                [
                    d.get("archivo") or d.get("asunto") or "",
                    d.get("error") or "",
                    d.get("factura") or "",
                    format_fecha_slash(d.get("fecha") or ""),
                    d.get("placa") or "",
                ]
            )
        for col in range(1, 6):
            ws_e.column_dimensions[get_column_letter(col)].width = 22

    widths = [12, 14, 12, 12, 16, 14, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_pdf(payload: dict) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(letter),
        leftMargin=1.2 * cm,
        rightMargin=1.2 * cm,
        topMargin=1.2 * cm,
        bottomMargin=1.2 * cm,
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "t",
        parent=styles["Heading2"],
        fontSize=13,
        spaceAfter=6,
    )
    small = ParagraphStyle(
        "s",
        parent=styles["Normal"],
        fontSize=8,
        leading=10,
    )
    cell = ParagraphStyle(
        "c",
        parent=styles["Normal"],
        fontSize=8,
        leading=10,
    )

    story = [
        Paragraph("aplicacionjunior · Servicio Trailero · Tanqueos", title),
        Paragraph(_meta(payload), small),
        Paragraph(
            f"Tanqueos: {payload.get('tanqueos_total', 0)} · "
            f"Errores: {payload.get('pdfs_error', 0)} · "
            f"Total: {_money(payload.get('total_general'))}",
            small,
        ),
        Spacer(1, 0.35 * cm),
    ]

    data = [
        [
            _pdf_cell("#", cell),
            _pdf_cell("Fecha", cell),
            _pdf_cell("Placa", cell),
            _pdf_cell("Cantidad", cell),
            _pdf_cell("Precio x Galón", cell),
            _pdf_cell("Descuento", cell),
            _pdf_cell("Valor Total", cell),
        ]
    ]
    for d in _rows_ok(payload):
        data.append(
            [
                _pdf_cell(str(d.get("tanqueo") or ""), cell),
                _pdf_cell(format_fecha_slash(d.get("fecha") or ""), cell),
                _pdf_cell(d.get("placa") or "", cell),
                _pdf_cell(_num(d.get("cantidad")), cell),
                _pdf_cell(_money(d.get("precio_galon")), cell),
                _pdf_cell(_money(d.get("descuento")), cell),
                _pdf_cell(_money(d.get("valor_total")), cell),
            ]
        )

    table = Table(
        data,
        colWidths=[1.2 * cm, 2.6 * cm, 2.4 * cm, 2.4 * cm, 3.2 * cm, 3 * cm, 3.2 * cm],
        repeatRows=1,
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D8D8E8")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#F5F3FF")],
                ),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story.append(table)

    bad = _rows_bad(payload)
    if bad:
        story.append(Spacer(1, 0.5 * cm))
        story.append(Paragraph("No leídos / con error", title))
        edata = [[_pdf_cell("Archivo / Asunto", cell), _pdf_cell("Error", cell)]]
        for d in bad:
            edata.append(
                [
                    _pdf_cell(d.get("archivo") or d.get("asunto") or "—", cell),
                    _pdf_cell(d.get("error") or "Error", cell),
                ]
            )
        et = Table(edata, colWidths=[10 * cm, 12 * cm], repeatRows=1)
        et.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DC2626")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#FECACA")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [colors.white, colors.HexColor("#FEF2F2")],
                    ),
                ]
            )
        )
        story.append(et)

    doc.build(story)
    return buf.getvalue()
