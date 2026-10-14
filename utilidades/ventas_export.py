"""Exporta listado de ventas a TXT, Excel y PDF."""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from effi_ventas import CUENTA_BANCOLOMBIA_ESPERADA


def _pdf_cell(text: str, style: ParagraphStyle) -> Paragraph:
    safe = (
        str(text or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )
    return Paragraph(safe, style)


def _money(n: float | int | None) -> str:
    try:
        return f"${float(n or 0):,.0f}".replace(",", ".")
    except (TypeError, ValueError):
        return "$0"


def _medio_label(d: dict) -> str:
    nombre = (d.get("medio_nombre") or "").strip()
    pasarela = (d.get("pasarela") or "").strip()
    if not pasarela and d.get("medio_key") == "pago_electronico":
        pasarela = "Mercado Pago"
    if nombre and pasarela:
        return f"{nombre} · {pasarela}"
    return nombre or pasarela or ""


def _meta(payload: dict) -> str:
    return (
        f"{payload.get('sucursal_nombre', '')} ({payload.get('codigo', '')}) · "
        f"{payload.get('desde', '')} → {payload.get('hasta', '')}"
    )


def _bloques(payload: dict) -> list[dict]:
    bloques = payload.get("por_sucursal") or []
    if payload.get("codigo") == "TODOS" and bloques:
        return bloques
    return []


def _detalle_por_sucursal(payload: dict, codigo: str) -> list[dict]:
    return [d for d in (payload.get("detalle") or []) if d.get("sucursal") == codigo]


def export_txt(payload: dict) -> bytes:
    lines = [
        "APLICACIONJUNIOR · LISTADO DE VENTAS",
        _meta(payload),
        f"Cuenta Bancolombia esperada: {payload.get('cuenta_esperada', CUENTA_BANCOLOMBIA_ESPERADA)}",
        f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        "",
    ]
    bloques = _bloques(payload)
    if bloques:
        for b in bloques:
            lines.append(f"=== {b.get('codigo')} · {b.get('nombre')} ===")
            for m in b.get("medios") or []:
                lines.append(
                    f"- {m.get('nombre')}: {m.get('facturas', 0)} factura(s) · "
                    f"{_money(m.get('total'))}"
                )
            lines.append(
                f"TOTAL {b.get('nombre')}: {b.get('facturas_total', 0)} factura(s) · "
                f"{_money(b.get('total_general'))}"
            )
            lines.append("")
        lines.append(
            f"TOTAL COMPLETO: {payload.get('facturas_total', 0)} factura(s) · "
            f"{_money(payload.get('total_general'))}"
        )
    else:
        lines.append("=== RESUMEN POR MEDIO ===")
        for m in payload.get("medios") or []:
            lines.append(
                f"- {m.get('nombre')}: {m.get('facturas', 0)} factura(s) · "
                f"{_money(m.get('total'))}"
            )
        lines.append(f"TOTAL GENERAL: {_money(payload.get('total_general'))}")

    lines.append("")
    lines.append("=== LISTADO ===")
    lines.append(
        f"{'Factura':<14} {'Num. PDF':<14} {'Fecha':<20} {'Cliente':<24} "
        f"{'Medio':<24} {'Cuenta':<14} {'Total':>12}"
    )
    lines.append("-" * 140)

    if bloques:
        for b in bloques:
            lines.append("")
            lines.append(f"-- {b.get('codigo')} · {b.get('nombre')} --")
            for d in _detalle_por_sucursal(payload, b.get("codigo") or ""):
                cuenta = d.get("cuenta") or "-"
                lines.append(
                    f"{(d.get('factura') or '')[:14]:<14} "
                    f"{(d.get('numeracion') or '-')[:14]:<14} "
                    f"{(d.get('fecha') or '')[:20]:<20} "
                    f"{(d.get('cliente') or '')[:24]:<24} "
                    f"{_medio_label(d)[:24]:<24} "
                    f"{cuenta[:14]:<14} "
                    f"{_money(d.get('total')):>12}"
                )
            lines.append(
                f"TOTAL {b.get('nombre')}: {_money(b.get('total_general'))} "
                f"({b.get('facturas_total', 0)} factura(s))"
            )
        lines.append("")
        lines.append(
            f"TOTAL COMPLETO: {_money(payload.get('total_general'))} "
            f"({payload.get('facturas_total', 0)} factura(s))"
        )
    else:
        for d in payload.get("detalle") or []:
            cuenta = d.get("cuenta") or "-"
            lines.append(
                f"{(d.get('factura') or '')[:14]:<14} "
                f"{(d.get('numeracion') or '-')[:14]:<14} "
                f"{(d.get('fecha') or '')[:20]:<20} "
                f"{(d.get('cliente') or '')[:24]:<24} "
                f"{_medio_label(d)[:24]:<24} "
                f"{cuenta[:14]:<14} "
                f"{_money(d.get('total')):>12}"
            )

    return ("\n".join(lines) + "\n").encode("utf-8")


def export_xlsx(payload: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas"

    yellow = PatternFill("solid", fgColor="FFF59D")
    header_fill = PatternFill("solid", fgColor="667EEA")
    section_fill = PatternFill("solid", fgColor="EEF0FA")
    sub_fill = PatternFill("solid", fgColor="F7F7FC")
    total_fill = PatternFill("solid", fgColor="F3F0FF")
    header_font = Font(color="FFFFFF", bold=True)
    bold = Font(bold=True)

    ws["A1"] = "aplicacionjunior · Listado de ventas"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = _meta(payload)
    ws["A3"] = (
        f"Cuenta Bancolombia esperada: "
        f"{payload.get('cuenta_esperada', CUENTA_BANCOLOMBIA_ESPERADA)}"
    )

    row = 5
    bloques = _bloques(payload)
    if bloques:
        ws.cell(row, 1, "Resumen por sucursal").font = bold
        row += 1
        for b in bloques:
            ws.cell(row, 1, f"{b.get('codigo')} · {b.get('nombre')}").font = bold
            for col in range(1, 4):
                ws.cell(row, col).fill = section_fill
            row += 1
            for col, h in enumerate(("Medio de pago", "Facturas", "Total"), 1):
                cell = ws.cell(row, col, h)
                cell.fill = header_fill
                cell.font = header_font
            row += 1
            for m in b.get("medios") or []:
                ws.cell(row, 1, m.get("nombre"))
                ws.cell(row, 2, m.get("facturas", 0))
                ws.cell(row, 3, float(m.get("total") or 0))
                ws.cell(row, 3).number_format = '"$"#,##0'
                row += 1
            ws.cell(row, 1, f"Total {b.get('nombre')}").font = bold
            ws.cell(row, 2, b.get("facturas_total", 0)).font = bold
            ws.cell(row, 3, float(b.get("total_general") or 0)).font = bold
            ws.cell(row, 3).number_format = '"$"#,##0'
            for col in range(1, 4):
                ws.cell(row, col).fill = sub_fill
            row += 2
        ws.cell(row, 1, "Total completo").font = bold
        ws.cell(row, 2, payload.get("facturas_total", 0)).font = bold
        ws.cell(row, 3, float(payload.get("total_general") or 0)).font = bold
        ws.cell(row, 3).number_format = '"$"#,##0'
        for col in range(1, 4):
            ws.cell(row, col).fill = total_fill
        row += 2
    else:
        ws.cell(row, 1, "Resumen por medio").font = bold
        row += 1
        for col, h in enumerate(("Medio de pago", "Facturas", "Total"), 1):
            cell = ws.cell(row, col, h)
            cell.fill = header_fill
            cell.font = header_font
        row += 1
        for m in payload.get("medios") or []:
            ws.cell(row, 1, m.get("nombre"))
            ws.cell(row, 2, m.get("facturas", 0))
            ws.cell(row, 3, float(m.get("total") or 0))
            ws.cell(row, 3).number_format = '"$"#,##0'
            row += 1
        ws.cell(row, 1, "Total general").font = bold
        ws.cell(row, 2, payload.get("facturas_total", 0)).font = bold
        ws.cell(row, 3, float(payload.get("total_general") or 0)).font = bold
        ws.cell(row, 3).number_format = '"$"#,##0'
        row += 2

    ws.cell(row, 1, "Listado de facturas").font = bold
    row += 1
    headers = [
        "Factura",
        "Numeración PDF",
        "Fecha",
        "Cliente",
        "Medio de pago",
        "Banco",
        "Cuenta",
        "Total",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    row += 1

    def write_detalle_row(d: dict) -> None:
        nonlocal row
        values = [
            d.get("factura"),
            d.get("numeracion") or "",
            d.get("fecha"),
            d.get("cliente"),
            _medio_label(d),
            d.get("banco") or "",
            d.get("cuenta") or "",
            float(d.get("total") or 0),
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row, col, val)
            if d.get("cuenta_alerta"):
                cell.fill = yellow
        ws.cell(row, 8).number_format = '"$"#,##0'
        row += 1

    if bloques:
        for b in bloques:
            ws.cell(row, 1, f"{b.get('codigo')} · {b.get('nombre')}").font = bold
            for col in range(1, 9):
                ws.cell(row, col).fill = section_fill
            row += 1
            for d in _detalle_por_sucursal(payload, b.get("codigo") or ""):
                write_detalle_row(d)
            ws.cell(row, 1, f"Total {b.get('nombre')}").font = bold
            ws.cell(row, 7, f"{b.get('facturas_total', 0)} factura(s)")
            ws.cell(row, 8, float(b.get("total_general") or 0)).font = bold
            ws.cell(row, 8).number_format = '"$"#,##0'
            for col in range(1, 9):
                ws.cell(row, col).fill = sub_fill
            row += 1
        ws.cell(row, 1, "Total completo").font = bold
        ws.cell(row, 7, f"{payload.get('facturas_total', 0)} factura(s)")
        ws.cell(row, 8, float(payload.get("total_general") or 0)).font = bold
        ws.cell(row, 8).number_format = '"$"#,##0'
        for col in range(1, 9):
            ws.cell(row, col).fill = total_fill
    else:
        for d in payload.get("detalle") or []:
            write_detalle_row(d)

    widths = [14, 14, 20, 28, 28, 14, 14, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_pdf(payload: dict) -> bytes:
    buf = io.BytesIO()
    page = landscape(letter)
    doc = SimpleDocTemplate(
        buf,
        pagesize=page,
        leftMargin=1.2 * cm,
        rightMargin=1.2 * cm,
        topMargin=1.2 * cm,
        bottomMargin=1.2 * cm,
    )
    usable = page[0] - 2.4 * cm
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "TitleLocal",
        parent=styles["Heading1"],
        fontSize=14,
        textColor=colors.HexColor("#4c51bf"),
        spaceAfter=6,
    )
    normal = ParagraphStyle("NormalLocal", parent=styles["Normal"], fontSize=9, leading=12)
    small = ParagraphStyle("SmallLocal", parent=styles["Normal"], fontSize=8, leading=10)
    cell = ParagraphStyle(
        "CellLocal",
        parent=styles["Normal"],
        fontSize=8,
        leading=10,
        wordWrap="CJK",
    )
    cell_right = ParagraphStyle(
        "CellRightLocal",
        parent=cell,
        alignment=2,
    )
    head_cell = ParagraphStyle(
        "HeadCellLocal",
        parent=styles["Normal"],
        fontSize=8,
        leading=10,
        textColor=colors.white,
        fontName="Helvetica-Bold",
    )

    story = [
        Paragraph("aplicacionjunior · Listado de ventas", title),
        Paragraph(_meta(payload), normal),
        Paragraph(
            f"Cuenta Bancolombia esperada: "
            f"<b>{payload.get('cuenta_esperada', CUENTA_BANCOLOMBIA_ESPERADA)}</b>",
            normal,
        ),
        Spacer(1, 0.35 * cm),
    ]

    def add_resumen_table(rows: list[list], highlight_last: bool = True) -> None:
        t = Table(rows, colWidths=[usable * 0.55, usable * 0.2, usable * 0.25])
        cmds = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#667eea")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#d8d8e8")),
            ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]
        if highlight_last:
            cmds.extend(
                [
                    ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#f3f0ff")),
                    ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                ]
            )
        t.setStyle(TableStyle(cmds))
        story.append(t)

    bloques = _bloques(payload)
    if bloques:
        story.append(Paragraph("<b>Resumen por sucursal</b>", normal))
        story.append(Spacer(1, 0.15 * cm))
        for b in bloques:
            story.append(
                Paragraph(
                    f"<b>{b.get('codigo')} · {b.get('nombre')}</b>",
                    normal,
                )
            )
            story.append(Spacer(1, 0.1 * cm))
            rows = [["Medio de pago", "Facturas", "Total"]]
            for m in b.get("medios") or []:
                rows.append(
                    [
                        m.get("nombre") or "",
                        str(m.get("facturas", 0)),
                        _money(m.get("total")),
                    ]
                )
            rows.append(
                [
                    f"Total {b.get('nombre')}",
                    str(b.get("facturas_total", 0)),
                    _money(b.get("total_general")),
                ]
            )
            add_resumen_table(rows)
            story.append(Spacer(1, 0.25 * cm))
        add_resumen_table(
            [
                ["Concepto", "Facturas", "Total"],
                [
                    "Total completo",
                    str(payload.get("facturas_total", 0)),
                    _money(payload.get("total_general")),
                ],
            ]
        )
    else:
        story.append(Paragraph("<b>Resumen por medio</b>", normal))
        story.append(Spacer(1, 0.15 * cm))
        resumen_data = [["Medio de pago", "Facturas", "Total"]]
        for m in payload.get("medios") or []:
            resumen_data.append(
                [m.get("nombre") or "", str(m.get("facturas", 0)), _money(m.get("total"))]
            )
        resumen_data.append(
            [
                "Total general",
                str(payload.get("facturas_total", 0)),
                _money(payload.get("total_general")),
            ]
        )
        add_resumen_table(resumen_data)

    story.extend(
        [Spacer(1, 0.45 * cm), Paragraph("<b>Listado</b>", normal), Spacer(1, 0.15 * cm)]
    )

    detalle_data = [[
        _pdf_cell("Factura", head_cell),
        _pdf_cell("Num. PDF", head_cell),
        _pdf_cell("Fecha", head_cell),
        _pdf_cell("Cliente", head_cell),
        _pdf_cell("Medio", head_cell),
        _pdf_cell("Cuenta", head_cell),
        _pdf_cell("Total", head_cell),
    ]]
    alert_rows = []
    section_rows = []
    subtotal_rows = []

    def add_detalle_line(d: dict) -> None:
        cuenta = d.get("cuenta") or "-"
        if d.get("banco") and cuenta != "-":
            cuenta = f"{d.get('banco')} {cuenta}"
        detalle_data.append(
            [
                _pdf_cell(d.get("factura") or "", cell),
                _pdf_cell(d.get("numeracion") or "-", cell),
                _pdf_cell((d.get("fecha") or "")[:19], cell),
                _pdf_cell(d.get("cliente") or "", cell),
                _pdf_cell(_medio_label(d), cell),
                _pdf_cell(cuenta, cell),
                _pdf_cell(_money(d.get("total")), cell_right),
            ]
        )
        idx = len(detalle_data) - 1
        if d.get("cuenta_alerta"):
            alert_rows.append(idx)

    if bloques:
        for b in bloques:
            detalle_data.append(
                [
                    _pdf_cell(
                        f"<b>{b.get('codigo')} · {b.get('nombre')}</b>",
                        cell,
                    ),
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                ]
            )
            section_rows.append(len(detalle_data) - 1)
            for d in _detalle_por_sucursal(payload, b.get("codigo") or ""):
                add_detalle_line(d)
            detalle_data.append(
                [
                    _pdf_cell(f"<b>Total {b.get('nombre')}</b>", cell),
                    "",
                    "",
                    "",
                    "",
                    _pdf_cell(f"{b.get('facturas_total', 0)} factura(s)", cell),
                    _pdf_cell(_money(b.get("total_general")), cell_right),
                ]
            )
            subtotal_rows.append(len(detalle_data) - 1)
        detalle_data.append(
            [
                _pdf_cell("<b>Total completo</b>", cell),
                "",
                "",
                "",
                "",
                _pdf_cell(f"{payload.get('facturas_total', 0)} factura(s)", cell),
                _pdf_cell(_money(payload.get("total_general")), cell_right),
            ]
        )
        subtotal_rows.append(len(detalle_data) - 1)
    else:
        for d in payload.get("detalle") or []:
            add_detalle_line(d)

    t2 = Table(
        detalle_data,
        colWidths=[
            usable * 0.09,
            usable * 0.09,
            usable * 0.13,
            usable * 0.18,
            usable * 0.17,
            usable * 0.22,
            usable * 0.12,
        ],
        repeatRows=1,
    )
    style_cmds = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#667eea")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d8d8e8")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("SPAN", (0, 0), (0, 0)),  # no-op placeholder removed below
    ]
    # Remove accidental SPAN
    style_cmds = [c for c in style_cmds if c[0] != "SPAN"]
    for r in section_rows:
        style_cmds.append(("BACKGROUND", (0, r), (-1, r), colors.HexColor("#EEF0FA")))
        style_cmds.append(("SPAN", (0, r), (-1, r)))
    for r in subtotal_rows:
        style_cmds.append(("BACKGROUND", (0, r), (-1, r), colors.HexColor("#F3F0FF")))
        style_cmds.append(("FONTNAME", (0, r), (-1, r), "Helvetica-Bold"))
    for r in alert_rows:
        style_cmds.append(("BACKGROUND", (0, r), (-1, r), colors.HexColor("#FFF59D")))
    t2.setStyle(TableStyle(style_cmds))
    story.append(t2)
    story.append(Spacer(1, 0.3 * cm))
    story.append(
        Paragraph(
            "Filas en amarillo: transferencia con cuenta distinta a "
            f"{payload.get('cuenta_esperada', CUENTA_BANCOLOMBIA_ESPERADA)}.",
            small,
        )
    )

    doc.build(story)
    return buf.getvalue()
